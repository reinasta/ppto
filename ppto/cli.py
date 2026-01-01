from __future__ import annotations

import json
import logging
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Iterable, List, Optional

import typer
from PIL import Image, ImageDraw
from pptx import Presentation

from .emit_latex import emit as emit_latex
from .emit_typst import emit as emit_typst
from .ir import converter, serializer
from .ir.models import Document

app = typer.Typer(add_completion=False)
logger = logging.getLogger(__name__)


def _detect_soffice() -> Optional[Path]:
    path = shutil.which("soffice")
    return Path(path) if path else None


def _convert_with_soffice(pptx_path: Path, out_dir: Path, slides: Optional[List[int]]) -> List[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [
        "soffice",
        "--headless",
        "--convert-to",
        "png",
        "--outdir",
        str(out_dir),
        str(pptx_path),
    ]
    logger.debug("Running LibreOffice conversion: %s", " ".join(cmd))
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    images = sorted(out_dir.glob("*.png"))
    if slides:
        allowed = {f"-{i}" for i in slides}
        images = [img for img in images if any(token in img.stem for token in allowed)]
    return images


def _fallback_render(pptx_path: Path, out_dir: Path, dpi: int, slides: Optional[List[int]]) -> List[Path]:
    presentation = Presentation(str(pptx_path))
    out_dir.mkdir(parents=True, exist_ok=True)
    slide_paths: List[Path] = []
    width_px = converter._emu_to_px(presentation.slide_width, dpi)
    height_px = converter._emu_to_px(presentation.slide_height, dpi)
    selected = set(slides) if slides else None
    for idx, slide in enumerate(presentation.slides, start=1):
        if selected and idx not in selected:
            continue
        image = Image.new("RGB", (int(width_px), int(height_px)), color="white")
        draw = ImageDraw.Draw(image)
        y_cursor = 10
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text.replace("\r", "\n")
                draw.text((10, y_cursor), f"[{shape.name}] {text}", fill="black")
                y_cursor += 20
        slide_file = out_dir / f"slide-{idx}.png"
        image.save(slide_file, dpi=(dpi, dpi))
        slide_paths.append(slide_file)
    return slide_paths


def _write_snippets(
    target: str, images: Iterable[Path], out_dir: Path, document: Optional[Document] = None
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    if target == "latex":
        suffix = ".tex"
        lines = [
            r"% Per-slide background snippets",
            *[
                rf"\usebackgroundtemplate{{\includegraphics[width=\paperwidth]{{{img.as_posix()}}}}}"
                for img in images
            ],
        ]
    else:
        suffix = ".typ"
        width = height = None
        if document:
            width = (document.size.width / document.dpi) * 2.54
            height = (document.size.height / document.dpi) * 2.54
        lines = [
            "#per-slide backgrounds",
            *[
                (
                    f"#box(width: {width:.2f}cm, height: {height:.2f}cm, image(\"{img.as_posix()}\"))"
                    if width and height
                    else f"#box(width: 100%, height: 100%, image(\"{img.as_posix()}\"))"
                )
                for img in images
            ],
        ]
    snippet_file = out_dir / f"snippets{suffix}"
    snippet_file.write_text("\n".join(lines), encoding="utf-8")
    return snippet_file


def _save_ir(document: Document, out_dir: Path) -> Path:
    ir_path = out_dir / "ir.json"
    serializer.document_to_json(document, ir_path)
    return ir_path


def _blank_images(document: Document, out_dir: Path) -> List[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    width = int(document.size.width)
    height = int(document.size.height)
    paths: List[Path] = []
    for slide in document.slides:
        color = "white"
        if slide.background.color:
            color = slide.background.color.hex
        image = Image.new("RGB", (width, height), color=color)
        path = out_dir / f"slide-{slide.index}.png"
        image.save(path, dpi=(document.dpi, document.dpi))
        paths.append(path)
    return paths


@app.command()
def convert(
    pptx_path: Path = typer.Argument(..., help="Path to input .pptx file"),
    out: Path = typer.Option(Path("out"), "--out", "-o", help="Directory for outputs"),
    slides: Optional[List[int]] = typer.Option(
        None, help="Optional slide numbers to render (1-based)"
    ),
    dpi: int = typer.Option(150, help="Render DPI for raster outputs"),
    target: str = typer.Option(
        "latex", help="Snippet target", case_sensitive=False, rich_help_panel="Output"
    ),
    mode: str = typer.Option(
        "images",
        help="Output mode: images (default) or ir (write intermediate representation)",
        case_sensitive=False,
    ),
) -> None:
    """Render a PPTX to slide PNGs and template snippets or IR."""
    target = target.lower()
    mode = mode.lower()
    out.mkdir(parents=True, exist_ok=True)

    document = converter.presentation_to_ir(pptx_path, slides=slides, dpi=dpi)
    if mode == "ir":
        ir_path = _save_ir(document, out)
        typer.echo(f"Wrote IR: {ir_path}")
        return

    soffice_path = _detect_soffice()
    if soffice_path:
        try:
            images = _convert_with_soffice(pptx_path, out, slides)
        except subprocess.CalledProcessError as exc:
            logger.warning("LibreOffice conversion failed, falling back: %s", exc)
            images = _fallback_render(pptx_path, out, dpi, slides)
    else:
        images = _fallback_render(pptx_path, out, dpi, slides)

    snippet_file = _write_snippets(target, images, out, document=document)
    typer.echo(f"Wrote {len(images)} slide images to {out}")
    typer.echo(f"Snippet file: {snippet_file}")


@app.command()
def emit(
    ir: Path = typer.Option(..., "--ir", exists=True, help="Path to IR JSON"),
    target: str = typer.Option(
        "latex", "--target", "-t", case_sensitive=False, help="Emission target"
    ),
    out: Path = typer.Option(Path("out"), "--out", "-o", help="Output directory"),
    images: Optional[List[Path]] = typer.Option(
        None,
        "--images",
        "-i",
        help="Optional per-slide images matching the IR slides order",
    ),
) -> None:
    """Emit LaTeX or Typst snippets from an IR file."""
    document = serializer.load_document(ir)
    out.mkdir(parents=True, exist_ok=True)
    target = target.lower()
    if images:
        image_paths = images
    else:
        tmp = Path(tempfile.mkdtemp(prefix="ppto-images-"))
        image_paths = _blank_images(document, tmp)
    if target == "latex":
        created = emit_latex(document, out, image_paths)
    elif target == "typst":
        created = emit_typst(document, out, image_paths)
    else:
        raise typer.BadParameter("Target must be latex or typst")

    typer.echo(f"Created {len(created)} {target} snippets in {out}")


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    app()


if __name__ == "__main__":
    main()
