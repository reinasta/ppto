from __future__ import annotations

import logging
from pathlib import Path
from typing import Iterable, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import models

EMU_PER_INCH = 914400

logger = logging.getLogger(__name__)


def _emu_to_px(value: int, dpi: int) -> float:
    return (value / EMU_PER_INCH) * dpi


def _color_to_hex(color) -> Optional[str]:
    try:
        rgb = color.rgb
        if rgb:
            return f"#{rgb:06X}"
    except Exception:
        return None
    return None


def _theme_colors(prs: Presentation) -> List[models.Color]:
    try:
        scheme = prs.part.theme_part.theme.color_scheme
        return [
            models.Color(hex=f"#{c.rgb:06X}")
            for c in scheme if getattr(c, "rgb", None) is not None
        ]
    except Exception:
        logger.debug("Unable to read theme colors", exc_info=True)
        return []


def _extract_text_box(shape, dpi: int, z_index: int) -> models.TextBox:
    runs: List[models.TextRun] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            runs.append(
                models.TextRun(
                    text=run.text,
                    font_size_pt=float(run.font.size.pt) if run.font.size else None,
                    color=(
                        models.Color(hex=_color_to_hex(run.font.color))
                        if _color_to_hex(run.font.color)
                        else None
                    ),
                    bold=run.font.bold,
                    italic=run.font.italic,
                )
            )

    fill_color = None
    if shape.fill and shape.fill.type and shape.fill.fore_color:
        hex_color = _color_to_hex(shape.fill.fore_color)
        fill_color = models.Color(hex=hex_color) if hex_color else None

    return models.TextBox(
        position=models.Point(
            x=_emu_to_px(shape.left, dpi), y=_emu_to_px(shape.top, dpi)
        ),
        size=models.Size(
            width=_emu_to_px(shape.width, dpi), height=_emu_to_px(shape.height, dpi)
        ),
        runs=runs,
        fill_color=fill_color,
        z_index=z_index,
    )


def _extract_image(shape, dpi: int, z_index: int) -> models.ImageElement:
    source = getattr(shape.image, "filename", "embedded-image")
    return models.ImageElement(
        position=models.Point(
            x=_emu_to_px(shape.left, dpi), y=_emu_to_px(shape.top, dpi)
        ),
        size=models.Size(
            width=_emu_to_px(shape.width, dpi), height=_emu_to_px(shape.height, dpi)
        ),
        source=source,
        description=shape.name,
        z_index=z_index,
    )


def _background(slide, dpi: int) -> models.Background:
    fill = getattr(slide, "background", None)
    if fill and fill.fill and fill.fill.type and fill.fill.fore_color:
        hex_color = _color_to_hex(fill.fill.fore_color)
        if hex_color:
            return models.Background(color=models.Color(hex=hex_color))
    return models.Background()


def presentation_to_ir(
    input_path: Path, slides: Optional[Iterable[int]] = None, dpi: int = 150
) -> models.Document:
    prs = Presentation(str(input_path))
    width_px = _emu_to_px(prs.slide_width, dpi)
    height_px = _emu_to_px(prs.slide_height, dpi)
    selected = set(slides) if slides else None

    slide_models: List[models.Slide] = []
    for idx, slide in enumerate(prs.slides, start=1):
        if selected and idx not in selected:
            continue
        text_boxes: List[models.TextBox] = []
        images: List[models.ImageElement] = []
        for z_index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text_boxes.append(_extract_text_box(shape, dpi, z_index))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images.append(_extract_image(shape, dpi, z_index))
        slide_models.append(
            models.Slide(
                index=idx,
                background=_background(slide, dpi),
                text_boxes=text_boxes,
                images=images,
            )
        )

    return models.Document(
        size=models.Size(width=width_px, height=height_px),
        dpi=dpi,
        theme=models.Theme(colors=_theme_colors(prs)),
        slides=slide_models,
    )
